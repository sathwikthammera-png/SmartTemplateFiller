from PySide6.QtWidgets import (
    QWidget, QVBoxLayout, QHBoxLayout, QLabel, QPushButton, QFileDialog, QMessageBox, QFrame
)
from PySide6.QtCore import Qt, Signal
from docx import Document
from ui.shared_state import set_template
import os


class ActionCard(QFrame):
    """Clickable dashboard card with premium hover effects and description."""
    clicked = Signal()

    def __init__(self, title, description, icon_char, highlight_color):
        super().__init__()
        self.setFrameShape(QFrame.StyledPanel)
        self.setObjectName("ActionCard")
        self.setCursor(Qt.PointingHandCursor)
        self.highlight_color = highlight_color
        
        self.setStyleSheet(f"""
            QFrame#ActionCard {{
                background-color: #111827;
                border: 1px solid #374151;
                border-radius: 12px;
            }}
            QFrame#ActionCard:hover {{
                border: 1.5px solid {highlight_color};
                background-color: #182235;
            }}
        """)
        
        layout = QVBoxLayout(self)
        layout.setContentsMargins(24, 30, 24, 30)
        layout.setSpacing(12)
        layout.setAlignment(Qt.AlignCenter)

        # Icon Label
        self.icon_lbl = QLabel(icon_char)
        self.icon_lbl.setAlignment(Qt.AlignCenter)
        self.icon_lbl.setStyleSheet(f"font-size: 44px; color: {highlight_color}; background-color: transparent;")
        layout.addWidget(self.icon_lbl)

        # Title Label
        self.title_lbl = QLabel(title)
        self.title_lbl.setAlignment(Qt.AlignCenter)
        self.title_lbl.setStyleSheet("font-size: 18px; font-weight: bold; color: white; background-color: transparent;")
        layout.addWidget(self.title_lbl)

        # Description Label
        self.desc_lbl = QLabel(description)
        self.desc_lbl.setAlignment(Qt.AlignCenter)
        self.desc_lbl.setWordWrap(True)
        self.desc_lbl.setStyleSheet("font-size: 12px; color: #9ca3af; background-color: transparent;")
        layout.addWidget(self.desc_lbl)

    def mousePressEvent(self, event):
        if event.button() == Qt.LeftButton:
            self.clicked.emit()
            event.accept()
        else:
            super().mousePressEvent(event)


class StartupScreen(QWidget):
    """Initial screen offering two premium choices:
    1) Create a new document from scratch
    2) Create document from template (open existing .docx)
    """

    def __init__(self, stack, screen1, screen2, creation_screen=None):
        super().__init__()
        self.stack = stack
        self.screen1 = screen1
        self.screen2 = screen2
        self.creation_screen = creation_screen
        self.init_ui()

    def init_ui(self):
        main_layout = QVBoxLayout(self)
        main_layout.setAlignment(Qt.AlignCenter)
        main_layout.setContentsMargins(40, 40, 40, 40)
        main_layout.setSpacing(25)
        self.setStyleSheet("background-color: #090d16;")

        # Header Block
        header_layout = QVBoxLayout()
        header_layout.setSpacing(6)
        header_layout.setAlignment(Qt.AlignCenter)

        title = QLabel("Smart Template Filler")
        title.setObjectName("Title")
        title.setAlignment(Qt.AlignCenter)
        title.setStyleSheet("color: white; font-size: 32px; font-weight: 800; background-color: transparent; letter-spacing: 0.5px;")
        header_layout.addWidget(title)

        subtitle = QLabel("AI-Powered Legal Document drafter & automation manager")
        subtitle.setAlignment(Qt.AlignCenter)
        subtitle.setStyleSheet("color: #9ca3af; font-size: 14px; background-color: transparent;")
        header_layout.addWidget(subtitle)
        main_layout.addLayout(header_layout)

        # Cards Layout (Side-by-Side)
        cards_layout = QHBoxLayout()
        cards_layout.setSpacing(24)
        cards_layout.setAlignment(Qt.AlignCenter)

        card_new = ActionCard(
            "Draft New Document",
            "Build a document from scratch. Access visual sections, dynamic formatting tables, numbered clauses, and design placeholders.",
            "➕",
            "#3b82f6"  # Blue Highlight
        )
        card_new.setFixedSize(280, 240)
        card_new.clicked.connect(self.create_new_document)
        cards_layout.addWidget(card_new)

        card_template = ActionCard(
            "Fill Word Template",
            "Load an existing Word file (.docx). Scan all document placeholders, fill in variables, and draft legal sections using our Questionnaire.",
            "📂",
            "#a855f7"  # Purple Highlight
        )
        card_template.setFixedSize(280, 240)
        card_template.clicked.connect(self.create_from_template)
        cards_layout.addWidget(card_template)

        main_layout.addLayout(cards_layout)

        # Footer Actions Group
        footer_layout = QVBoxLayout()
        footer_layout.setSpacing(12)
        footer_layout.setAlignment(Qt.AlignCenter)

        btn_open_template_screen = QPushButton("Advanced Drag & Drop Upload...")
        btn_open_template_screen.setFixedWidth(260)
        btn_open_template_screen.setStyleSheet("""
            QPushButton {
                background-color: transparent;
                color: #9ca3af;
                border: 1px solid #374151;
                border-radius: 6px;
                padding: 10px 16px;
                font-weight: 500;
            }
            QPushButton:hover {
                background-color: #111827;
                color: white;
                border-color: #4b5563;
            }
        """)
        btn_open_template_screen.clicked.connect(self.open_template_screen)
        footer_layout.addWidget(btn_open_template_screen)

        note = QLabel("Smart Template Filler • Pro Suite")
        note.setAlignment(Qt.AlignCenter)
        note.setStyleSheet("color: #4b5563; font-size: 11px; font-weight: bold; background-color: transparent;")
        footer_layout.addWidget(note)

        main_layout.addLayout(footer_layout)
        self.setLayout(main_layout)

    def create_new_document(self):
        if self.creation_screen:
            self.stack.setCurrentWidget(self.creation_screen)
            return

        try:
            doc = Document()
            set_template("", doc)
            try:
                self.screen2.build_form_fields()
            except Exception:
                pass
            self.stack.setCurrentWidget(self.screen2)
        except Exception as e:
            QMessageBox.critical(self, "Error", f"Failed to create document:\n{e}")

    def create_from_template(self):
        filepath, _ = QFileDialog.getOpenFileName(self, "Select Template", "", "Word Documents (*.docx)")
        if not filepath:
            return
        try:
            doc = Document(filepath)
            set_template(filepath, doc)
            try:
                self.screen2.build_form_fields()
            except Exception:
                pass
            self.stack.setCurrentWidget(self.screen2)
        except Exception as e:
            QMessageBox.critical(self, "Error", f"Failed to load template:\n{e}")

    def open_template_screen(self):
        try:
            self.stack.setCurrentWidget(self.screen1)
        except Exception as e:
            QMessageBox.critical(self, "Error", f"Failed to open template screen:\n{e}")
